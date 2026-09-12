import collections
import datetime as dt
import json
import re
import warnings
import openpyxl

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
path = 'Накопичення.xlsx'
wb = openpyxl.load_workbook(path, data_only=False)
cached = openpyxl.load_workbook(path, data_only=True)
out = []
for ws in wb:
    cs = cached[ws.title]
    cells = [c for row in ws for c in row if c.value is not None]
    formulas = [c for c in cells if c.data_type == 'f']
    dates = [(c.coordinate, c.value.date()) for c in cells if isinstance(c.value, dt.datetime)]
    text_dates = [(c.coordinate,c.value) for c in cells if isinstance(c.value,str) and re.fullmatch(r'\d{1,2}[./-]\d{1,2}[./-]\d{4}',c.value.strip())]
    funcs = collections.Counter(f.upper() for c in formulas for f in re.findall(r'([A-Za-z_.]+)\s*\(',c.value))
    errors = collections.Counter(str(c.value) for row in cs for c in row if c.data_type=='e')
    headers = [(c.coordinate, c.value) for row in ws.iter_rows(min_row=1,max_row=2 if ws.title=='ОВгП' else 1) for c in row if c.value is not None]
    missing_cache = [c.coordinate for c in formulas if cs[c.coordinate].value is None]
    formula_samples = {}
    for c in formulas:
        for f in re.findall(r'([A-Za-z_.]+)\s*\(',c.value):
            if f.upper() not in formula_samples:
                formula_samples[f.upper()] = c.coordinate
    # Only data-quality metadata, no operational values or comparisons.
    detail = {
        'sheet':ws.title, 'extent':ws.calculate_dimension(),
        'nonempty_rows':len(set(c.row for c in cells)), 'nonempty_cells':len(cells),
        'headers':headers, 'tables':[{ 'name':t.name,'range':t.ref} for t in ws.tables.values()],
        'merged_count':len(ws.merged_cells.ranges),
        'formula_cells':len(formulas), 'functions':funcs, 'formula_examples_at':formula_samples,
        'formula_missing_cache':len(missing_cache), 'missing_cache_examples':missing_cache[:5],
        'cached_errors':errors,
        'date_cells':len(dates), 'distinct_dates':len(set(v for _,v in dates)),
        'min_date':min((v for _,v in dates),default=None), 'max_date':max((v for _,v in dates),default=None),
        'text_date_count':len(text_dates), 'text_date_cells':[a for a,v in text_dates[:8]],
        'empty_header_columns':[openpyxl.utils.get_column_letter(col) for col in range(1,ws.max_column+1) if ws.cell(1,col).value is None],
        'hidden':ws.sheet_state,
    }
    if ws.title in ('ГОЧ','Втрати ОС','Аркуш3','Застосування БК та FPV'):
        body=[r for r in range(2,ws.max_row+1) if any(ws.cell(r,c).value is not None for c in range(1,ws.max_column+1))]
        key_cols=1 if ws.title=='Застосування БК та FPV' else 2
        keys=collections.Counter(tuple(str(cs.cell(r,c).value) for c in range(1,key_cols+1)) for r in body)
        detail['duplicate_key_groups']=sum(n>1 for n in keys.values())
        detail['duplicate_key_extra_rows']=sum(n-1 for n in keys.values() if n>1)
        detail['rows_missing_date']=sum(cs.cell(r,1).value is None for r in body)
        detail['rows_missing_group']=sum(cs.cell(r,2).value is None for r in body) if key_cols==2 else None
        detail['data_columns']=[]
        for col in range(key_cols+1,ws.max_column+1):
            counts=collections.Counter('blank' if cs.cell(r,col).value is None else 'error' if cs.cell(r,col).data_type=='e' else 'number' if isinstance(cs.cell(r,col).value,(int,float)) else 'text' for r in body)
            detail['data_columns'].append({'column':openpyxl.utils.get_column_letter(col),'types':counts})
    out.append(detail)
print(json.dumps(out,ensure_ascii=False,default=str,indent=2))
